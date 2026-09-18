#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把 docs/09 的技术报告草稿写进官方《作品提交模板》.docx。

用法：
    python3 tools/phywear/make_report_docx.py \
        --template "/home/xpqh/下载/2026 首届 openvela AI 硬件开发者大赛 - 作品提交模板.docx" \
        --out docs/PhyWear_技术报告_官方模板.docx

做法：**打开官方模板本体**（保留其样式与格式），在对应标题段落后插入草稿内容，
并把「1、信息表」的空表格填上；不新建模板、不改动其它小节标题。
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

import docx
from docx.shared import Pt

def _find_draft() -> Path:
    """在 cwd 与脚本上级目录里找草稿（脚本可能放在 tools/ 或 tools/phywear/）。"""
    name = "09_官方提交模板_填写草稿.md"
    for base in [Path.cwd(), *Path(__file__).resolve().parents]:
        cand = base / "docs" / name
        if cand.is_file():
            return cand
    raise SystemExit("找不到 docs/09_官方提交模板_填写草稿.md（请在参赛仓根目录运行本脚本）")


DRAFT = _find_draft()

# 官方模板里的标题 → 草稿里的起始标记
SECTION_MAP = [
    ("3.1", "## 3.1 绪论"),
    ("3.2", "## 3.2 系统方案设计"),
    ("3.3", "## 3.3 核心算法与技术原理"),
    ("3.4", "## 3.4 系统实现"),
    ("3.5", "## 3.5 系统测试与结果分析（量化）"),
    ("3.6", "## 3.6 AI-Native 开发说明"),
    ("3.7", "## 3.7 总结与展望"),
]
INFO_ROWS = {
    "作品名称": "openvela 腕上智慧物理工坊（PhyWear）",
    "队伍名称": "contest2026_427_xinpingqihe（队长 XPQHyue，成员 程秋霞）",
    "团队分工": "XPQHyue：架构/驱动/系统适配/AI Agent 集成/文档与证据；程秋霞：界面与交互实现、演示视频拍摄与剪辑",
    "选题方向": "③ 新硬件平台适配（主线）；① AI 硬件产品创新：已交付「主动 + 执行」场景并真机验证",
}


def parse_draft(text: str) -> dict[str, list[str]]:
    """把草稿按 '## x.y' 切成 {标记: [行]}（去掉 markdown 强调符号）。"""
    out: dict[str, list[str]] = {}
    cur = None
    for raw in text.splitlines():
        line = raw.rstrip()
        if line.startswith("## "):
            cur = line
            out[cur] = []
        elif cur is not None:
            clean = re.sub(r"\*\*(.+?)\*\*", r"\1", line)
            out[cur].append(clean)
    return out


def fill_info_table(doc) -> int:
    filled = 0
    for table in doc.tables:
        for row in table.rows:
            if len(row.cells) < 2:
                continue
            key = row.cells[0].text.strip()
            cur = row.cells[1].text.strip()
            # 空单元格，或还是模板占位文字（如"各成员姓名与承担的工作"），都要覆盖
            placeholder = cur == "" or "各成员姓名" in cur or "可多选组合" in cur or cur.startswith("xx")
            for k, v in INFO_ROWS.items():
                if k in key and placeholder:
                    row.cells[1].text = v
                    filled += 1
    return filled


def add_md_table(doc, anchor, block):
    """把一段 markdown 表格行（| a | b |）转成真正的 docx 表格，并插到 anchor 之后。"""
    rows = []
    for r in block:
        cells = [c.strip() for c in r.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c or "-") for c in cells):
            continue                      # 分隔行
        rows.append(cells)
    if not rows:
        return anchor
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol)
    try:
        t.style = "Table Grid"            # 模板可能没有该样式，拿不到就退回默认
    except Exception:
        pass
    for i, r in enumerate(rows):
        for j in range(ncol):
            t.cell(i, j).text = r[j] if j < len(r) else ""
            if i == 0:
                for para in t.cell(i, j).paragraphs:
                    for run in para.runs:
                        run.bold = True
    anchor.addnext(t._tbl)
    return t._tbl


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--draft", default=str(DRAFT))
    args = ap.parse_args()

    draft = parse_draft(Path(args.draft).read_text(encoding="utf-8"))
    doc = docx.Document(args.template)

    filled = fill_info_table(doc)
    inserted = 0
    for para in doc.paragraphs:
        text = para.text.strip()
        for _sec, marker in SECTION_MAP:
            if text.startswith(marker.replace("## ", "")) or text.startswith(_sec):
                body = [l for l in draft.get(marker, []) if l.strip()]
                if not body:
                    continue
                anchor = para._p
                i = 0
                while i < len(body):
                    line = body[i]
                    if not line.strip():
                        i += 1
                        continue
                    if line.lstrip().startswith("|"):
                        blk = []
                        while i < len(body) and body[i].lstrip().startswith("|"):
                            blk.append(body[i]); i += 1
                        anchor = add_md_table(doc, anchor, blk)
                        inserted += 1
                        continue
                    new = doc.add_paragraph()
                    new.paragraph_format.space_after = Pt(2)
                    new.add_run(line.strip())
                    anchor.addnext(new._p)
                    anchor = new._p
                    inserted += 1
                    i += 1
                break

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(f"✅ 已生成 {out}（信息表填入 {filled} 格，正文插入 {inserted} 段）")
    print("⚠️ 请打开 .docx 目检：删掉每节自带的提示文字（如'xx%（可估算…）'占位），确认无空节。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
